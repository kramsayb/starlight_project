from openpyxl import Workbook

from src.data.models import Batch
from src.utils.datetime_utils import report_date

def excel_generator(batch: Batch, stats: dict, file_path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Информация о партии"
    ws2 = wb.create_sheet("Продукция")
    ws3 = wb.create_sheet("Статистика")
    ws.append(["Номер партии", batch.batch_number])
    ws.append(["Дата партии", batch.batch_date])
    ws.append(["Статус", "Закрыта" if batch.is_closed else "Открыта"])
    ws.append(["Рабочий центр", batch.work_center.name])
    ws.append(["Смена", batch.shift])
    ws.append(["Бригада", batch.team])
    ws.append(["Номенклатура", batch.nomenclature])
    ws.append(["Начало смены", report_date(batch.shift_start)])
    ws.append(["Окончание смены",  report_date(batch.shift_end)])
    ws2.append(["ID", "Уникальный код", "Агрегирована", "Дата агрегации"])
    for product in batch.products:
        ws2.append([
            product.id,
            product.unique_code,
            'Да' if product.is_aggregated else 'Нет',
            report_date(product.aggregated_at) if product.aggregated_at else '-'
        ])
    ws3.append(["Всего продукции", stats["total"]])
    ws3.append(["Агрегировано", stats["aggregated"]])
    ws3.append(["Осталось", stats["remaining"]])
    ws3.append(["Процент выполнения", f"{stats['percent']}%"])
    ws3.append(["Средняя скорость", f"{stats['rate_per_hour']} ед/час"])
    wb.save(file_path)