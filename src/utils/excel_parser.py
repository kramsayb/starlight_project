import os
from typing import List, Dict

from openpyxl import load_workbook
from csv import DictReader


def parse_excel(file_path: str) -> List[Dict]:
    wb = load_workbook(file_path, read_only=True)
    try:
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        headers = next(rows)
        result = []
        for row_number, row in enumerate(rows, start=2):
            if all(value is None for value in row):
                continue
            item = dict(zip(headers, row))
            item["_row"] = row_number
            result.append(item)
    finally:
        wb.close()
    return result


def parse_csv(file_path: str) -> List[Dict]:
    result = []
    with open(file_path, encoding="utf-8-sig", newline="") as file:
        reader = DictReader(file, delimiter=";")
        for row in reader:
            if all(not value for value in row.values()):
                continue
            row["_row"] = reader.line_num
            result.append(row)
    return result


def parse_file(file_path: str) -> List[Dict]:
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".xlsx":
        return parse_excel(file_path)
    elif ext == ".csv":
        return parse_csv(file_path)
    else:
        raise ValueError(f"Unsupported file type: {ext}")
